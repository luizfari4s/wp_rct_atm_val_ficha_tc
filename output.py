from process import logger
from pandas import Categorical

from process import obter_periodo_produtivo
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
def arquivos_de_saida(df_rct_batch_vivos_pend_importacao,df_ficha,df_periodo,pst_egj,pst_out):
    logger.write( etapa='output', mensagem=f'ref identificacao fichas validadas')
    
    df_ficha['PanelSmart#1'] = df_ficha['PanelSmart#1'].astype(str)
    ficha_out = df_ficha.merge(df_rct_batch_vivos_pend_importacao[['Validacao_Origem','Duplicidade_de_Email','Duplicidade_de_Telefone','flag_complementar',
                                                                'saida_decisao','PanelSmart#1','Pre_Origen','Origen_Recrutado_32_IHS','Origen_Recrutado_42_Expansao',
                                                                'Origem_importar','Decisao_Final','FaixaNumCompras','validacao_gacode','dt_ficha','mes_produtivo_ficha',
                                                                'Data_Entrada_GPM','mes_produtivo',
                                                                'Lote','Data_inicio_importacao','data_processamento','gpm','Duplicidade_de_Email',
                                                                    'Duplicidade_de_Telefone', 'Classific', 'chave_setor','FSPanel1'  #Base LOTE
                                                                ]], on='PanelSmart#1', how='left')

    logger.write( etapa='output', mensagem=f'ref regra hard code linha 18 (output.py) | lotes para não importar')
    
    from datetime import datetime
    ficha_out.loc[(ficha_out.Decisao_Final.isin(['SUBIR GPM - CRITERIO DE QUALIDADE (<25 ATOS)',
                                                'SUBIR GPM - AJUSTAR ORIGEM',
                                                'SUBIR GPM'])) & (ficha_out.gpm == True), 'Lote'] = 'NAO_IMPORTAR_' + str(datetime.now().strftime('%Y%m%d_%H%M%S'))


    logger.write( etapa='output', mensagem=f'ref separação de fichas')
    ficha_out = ficha_out[[
        'data_processamento','gpm','Status','Lote','Data_inicio_importacao',
        
        'dt_ficha','mes_produtivo_ficha','flag_complementar','Pre_Origen','Origen_Recrutado_32_IHS','Origen_Recrutado_42_Expansao','Origem_importar','validacao_gacode','Validacao_Origem',
        'Data_Entrada_GPM','mes_produtivo',
        'saida_decisao','FaixaNumCompras','Decisao_Final',
        
        'Gacode',
        
        'PanelSmart#1','Estado','Entrevistador#1','Entrevistador#2','Entrevistador#3','Cota_Origem','Cidade_Cota',
        'P2','P7','P8','P9_3#1','P9_3#5','P9_3#6','P9_3#9','P9_3#10','P9_3#12','P9_3#22','P9_3#24','P9_3#25','P9_3#49',
        'P9_3#59','P9_3#62','P9_3#66','P9_3#70','P9_3#71','P9_3#72','P9_3#73','P9_1#68','P9_1#69','P10','P10a#1','P10b#1',
        'P10c','P10d_t','P10e#1_t','P10e#2_t','P10e#3_t','P10f','P10g','P10h','P10i#1','P10k','P12B#1','P12B#2','P12B#3',
        'P12B#4','P12B#5','P12B#6','P12B#7','P12B#8','P12B#9','P12B#10','P12B_1#1','P12B_1#2','P12B_1#3','P12B_1#4','P12B_1#5',
        'P12B_1#6','P12B_1#7','P12B_1#8','P12B_1#9','P12B_1#10','P12C_1','P12C_2','P12C_3','P12C_4','P12C_5','P12C_6','P12C_7',
        'P12C_8','P12C_9','P12C_10','P12D#1','P12D#2','P12D#3','P12D#4','P12D#5','P12D#6','P12D#7','P12D#8','P12D#9','P12D#10',
        'P12E#1','P12E#2','P12E#3','P12E#4','P12E#5','P12E#6','P12E#7','P12E#8','P12E#9','P12E#10','P12G#1_t','P12G#2_t','P12G#3_t',
        'P12G#4_t','P12G#5_t','P12G#6_t','P12G#7_t','P12G#8_t','P12G#9_t','P12G#10_t','P12H#1_t','P12H#2_t','P12H#3_t','P12H#4_t',
        'P12H#5_t','P12H#6_t','P12H#7_t', 'P12H#8_t','P12H#9_t','P12H#10_t','P12I#1','P12I#2','P12I#3','P12I#4','P12I#5','P12I#6',
        'P12I#7','P12I#8','P12I#9','P12I#10','P13b','P13c','P13','P13a',
        'P14','P14a#1','P14a#2','P14a#3','P14a#4','P14a#5','P14a#6','P14a#7','P14b#1','P14b#2','P14b#3','P14b#4','P14b#5','P14b#6','P14b#7',

        'P14c#1','P14c#2','P14c#3','P14c#4','P14c#5',
        'P14c#6','P14c#7','P14d#1','P14d#2','P14d#3','P14d#4','P14d#5','P14d#6','P14d#7','PF14e#1','PF14e#2','PF14e#3','PF14e#4',
        'PF14e#5','PF14e#6','PF14e#7','PF14f#1','PF14f#2','PF14f#3','PF14f#4','PF14f#5','PF14f#6','PF14f#7', 'AVATAR FINALIZADO (BRT)','P12a#1','Duplicidade_de_Email',
        'Duplicidade_de_Telefone', 'Classe', 'Classific', 'chave_setor','FSPanel1'


    ]]
    logger.write( etapa='output', mensagem=f'ref atualização detalhe validacao')

    ficha_out.rename(columns={'P12a#1':'qtd_indiv(P12a#1)',
                                                        'data_processamento':'data_processamento'}, inplace=True)


    ficha_out.loc[~ficha_out['AVATAR FINALIZADO (BRT)'].isna(), 'FICHA_TC'] = 'NO_GPM'
    ficha_out.loc[ficha_out['AVATAR FINALIZADO (BRT)'].isna(), 'FICHA_TC'] = 'FA_FICHA'


    detalhe = ficha_out[[
        'Lote','Data_inicio_importacao', 'Status',# Base LOTES
        'dt_ficha',
        'mes_produtivo_ficha',
        'Decisao_Final',
        'PanelSmart#1',
        'qtd_indiv(P12a#1)',
        'Classe',
        'Classific',
        'FaixaNumCompras',
        'Data_Entrada_GPM',
        'FSPanel1',
        'mes_produtivo',
        'chave_setor',
        'Gacode',
        'validacao_gacode',
        'Pre_Origen',
        'Origen_Recrutado_32_IHS',
        'Origen_Recrutado_42_Expansao',
        'Origem_importar',
        'Duplicidade_de_Email',
        'Duplicidade_de_Telefone', 
        'Validacao_Origem',
        'saida_decisao',
        'flag_complementar',
        'data_processamento'
        
        
    ]]

    detalhe.loc[detalhe.mes_produtivo == '2026-08'].Status.value_counts(dropna=False)

    detalhe.drop_duplicates(inplace=True)

    logger.write( etapa='output', mensagem=f'ref atualização overview')

    overview = (
        detalhe
        .groupby('Decisao_Final')['PanelSmart#1']
        .nunique()
        .reset_index()
    ).rename(columns={'PanelSmart#1':'iddomicilio'})

    ordem = [
        'OFF - TESTE',
        'OFF - MORTALIDADE',
        'SUBIR GPM',
        'OFF - ABAIXO DE 5 ATOS',
        'SUBIR GPM - AJUSTAR ORIGEM',
        'OFF - REGIÃO FORA DA COLETA',
        'SUBIR GPM - ACIMA DE 5 ATOS)',
        'GPM - FICHA IMPORTADA'
    ]

    overview['Decisao_Final'] = Categorical(
        overview['Decisao_Final'],
        categories=ordem,
        ordered=True
    )

    overview = overview.sort_values('Decisao_Final').reset_index(drop=True)

        
    logger.write( etapa='output', mensagem=f'ref atualização bolsa engajamento')


    leads_bolsa = ficha_out[
     (ficha_out.gpm == False) &      
    ~(ficha_out.Lote.isna()) &
     (ficha_out.Decisao_Final.isin(['SUBIR GPM - CRITERIO DE QUALIDADE (<25 ATOS)',
                                    'SUBIR GPM - AJUSTAR ORIGEM',
                                    'SUBIR GPM']))][[
    "P10b#1",
    'P10a#1',
    'data_processamento',
    'Status',
    'Lote',
    'Data_inicio_importacao',
    'dt_ficha',
    'mes_produtivo_ficha',
    'Data_Entrada_GPM',
    'mes_produtivo',
    'Origem_importar',''
    'FaixaNumCompras',
    'Decisao_Final',
    'PanelSmart#1',
    'P10e#2_t',
    'P10d_t',
    'Estado',
    'Entrevistador#1',
    'Entrevistador#2',
    'gpm',
    ]]

    leads_bolsa['Nome'] = leads_bolsa["P10a#1"]+ " "+leads_bolsa["P10b#1"]

    leads_bolsa.drop(columns=['P10a#1','P10b#1'],inplace=True)


    leads_bolsa.rename(columns={
        'Data_Entrada_GPM':'dt_entrada_GPM',
        'mes_produtivo':'mes_produtivo_GPM',
        'PanelSmart#1':'IdDomicilio',
        'P10e#2_t':'Telefone',
        'P10d_t': 'Email',
        'Entrevistador#1':'Cidade',
        'Entrevistador#2':'Bairro'
    },inplace=True)

    leads_bolsa= leads_bolsa.loc[leads_bolsa.Lote.str.startswith('GPM')]

    data_atual = datetime.now()
    nm_arquivo = f'/BOLSA_ENGAJAMENTO_{data_atual.strftime('%d%m%Y')}.xlsx'
    output_bolsa = str(pst_egj) + nm_arquivo

    leads_bolsa.to_excel(output_bolsa, index=False)

    logger.write( etapa='output', mensagem=f'ref geração do arquivo final')


    def formatar_sheet(ws):

        # Cores
        cinza = PatternFill("solid", fgColor="CFCFCF")
        azul = PatternFill("solid", fgColor="E3E7FA")
        amarelo = PatternFill("solid", fgColor="FDFFBA")
        vermelho = PatternFill("solid", fgColor="FFD8C2")

        # Cabeçalho
        for cell in ws[1]:
            cell.fill = cinza
            cell.font = Font(color="000000", bold=True)

        # Congela primeira linha
        ws.freeze_panes = "A2"

        # Filtros
        ws.auto_filter.ref = ws.dimensions

        # Ajuste automático da largura
        for col in ws.columns:
            largura = max(len(str(c.value)) if c.value else 0 for c in col) + 2
            ws.column_dimensions[get_column_letter(col[0].column)].width = largura

        # Colunas que deseja destacar
        colunas_azuis = [
            'Validacao_Origem',
            'saida_decisao',
            'Decisao_Final',
            'Pre_Origen',
            '(PNC)Origen_Recrutado_32_IHS',
            '(EXPANSÃO)Origen_Recrutado_42_Expansao',
            'FaixaNumCompras'
        ]

        colunas_amarelas = [
            'Origem_importar',
            'validacao_gacode',
            'Gacode'
            
        ]
        
        colunas_avermelhadas = [
            'flag_complementar'
            
        ]

        for cell in ws[1]:
            if cell.value in colunas_azuis:
                for linha in range(2, ws.max_row + 1):
                    ws.cell(row=linha, column=cell.column).fill = azul

            elif cell.value in colunas_amarelas:
                for linha in range(2, ws.max_row + 1):
                    ws.cell(row=linha, column=cell.column).fill = amarelo
                    
            elif cell.value in colunas_avermelhadas:
                for linha in range(2, ws.max_row + 1):
                    ws.cell(row=linha, column=cell.column).fill = vermelho

    from datetime import datetime
    from pathlib import Path
    import pandas as pd
    
    periodo_real = obter_periodo_produtivo(data_atual, df_periodo)

    nm_arquivo = f'/PROD_VALIDACAO_FICHA_TC_{periodo_real}.xlsx'
    output = str(pst_out) + nm_arquivo

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        # Abas principais
        overview.to_excel(writer, sheet_name="Resumo", index=False)
        detalhe.to_excel(writer, sheet_name="Detalhe Validação", index=False)

        # Cria uma aba para cada Decisao_Final
        for decisao, df in ficha_out.groupby("Decisao_Final", sort=False):

            # Remove caracteres inválidos para nomes de abas
            nome_aba = (
                str(decisao)
                .replace("/", "-")
                .replace("\\", "-")
                .replace("*", "")
                .replace("?", "")
                .replace("[", "")
                .replace("]", "")
                .replace(":", "-")
            )[:31]  # Limite do Excel

            df.to_excel(writer, sheet_name=nome_aba, index=False)
        
        # Ajustando Cores de Colunas
        wb = writer.book

        for ws in wb.worksheets:
            formatar_sheet(ws)
            
        
    logger.write( etapa='output', mensagem=f"Arquivo '{output}' criado com sucesso!")



